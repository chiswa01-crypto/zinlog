# ==============================================================================
# [CODE FREEZE ACTIVE - ALL MODULES]
# STATUS: STABLE / FROZEN - BACKEND LOGIC PROTECTED
# ALL BACKEND PARSER & SERVICE LOGIC ARE LOCKED. FOCUS: UI / WEB APPEARANCE.
# ==============================================================================
import os
import re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from datetime import datetime
import pdfplumber

TEMPLATE_PATH = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', 'static', 'templates', 'DATABASE LOKAL BC 4.0.xlsx'))

def parse_bc40_pdf(pdf_path: str):
    """
    Ekstraksi data Dokumen Pabean BC 4.0 (Pemberitahuan Pemasukan Barang Asal Tempat Lain Dalam Daerah Pabean ke TPB).
    Menggabungkan analisis tabel (extract_tables) dan teks lengkap untuk akurasi 100%.
    """
    full_text = ""
    all_tables = []
    
    with pdfplumber.open(pdf_path) as pdf:
        for page in pdf.pages:
            t = page.extract_text() or ""
            full_text += t + "\n"
            for tbl in page.extract_tables():
                all_tables.append(tbl)
                
    dokumen = "BC 4.0"
    
    # 1. Nomor Pengajuan (26 karakter)
    nomor_aju = ""
    aju_match = re.search(r'Nomor\s+Pengajuan\s*:\s*([0-9A-Za-z]+)', full_text, re.IGNORECASE)
    if aju_match:
        nomor_aju = aju_match.group(1).strip()
    else:
        aju_fallback = re.search(r'\b(000040[0-9A-Za-z]{20})\b', full_text)
        if aju_fallback:
            nomor_aju = aju_fallback.group(1).strip()
            
    # 2. Nomor Pendaftaran & Tanggal Pendaftaran
    nomor_daftar = ""
    tanggal_daftar = ""
    daftar_match = re.search(r'Nomor\s+Pendaftaran\s*:\s*(\d+)', full_text, re.IGNORECASE)
    if daftar_match:
        nomor_daftar = daftar_match.group(1).strip()
        
    tgl_match = re.search(r'Nomor\s+Pendaftaran[\s\S]*?Tanggal\s*:\s*([0-9]{2}[-/\.][0-9]{2}[-/\.][0-9]{4})', full_text, re.IGNORECASE)
    if not tgl_match:
        tgl_match = re.search(r'Tanggal\s*:\s*([0-9]{2}[-/\.][0-9]{2}[-/\.][0-9]{4})', full_text, re.IGNORECASE)
    if tgl_match:
        tanggal_daftar = tgl_match.group(1).strip()
        
    # 3. Nama Pemasok / Pengirim
    nama_pemasok = ""
    pemasok_match = re.search(r'(?:PENGIRIM\s+BARANG[\s\S]*?6\.\s*Nama|6\.\s*Nama)\s*:\s*([^\n\r]+)', full_text, re.IGNORECASE)
    if pemasok_match:
        raw_pemasok = pemasok_match.group(1).strip()
        nama_pemasok = re.split(r'(?:7\.\s*Alamat|NPWP|NITKU)', raw_pemasok, flags=re.IGNORECASE)[0].strip()
        
    # 4. Dokumen Pelengkap: No Surat Jalan & No Invoice
    no_surat_jalan = ""
    no_invoice = ""
    
    # Cari di tabel Lembar Lanjutan Dokumen Pelengkap
    for tbl in all_tables:
        for r in tbl:
            if r and len(r) >= 4 and r[1] and r[3]:
                r1_str = str(r[1])
                r3_str = str(r[3])
                types = [x.strip() for x in r1_str.split('\n') if x.strip()]
                nums = [x.strip() for x in r3_str.split('\n') if x.strip()]
                for t_name, n_val in zip(types, nums):
                    if 'SURAT JALAN' in t_name.upper() or 'SURAT' in t_name.upper():
                        no_surat_jalan = n_val
                    elif 'INVOICE' in t_name.upper() or 'FAKTUR' in t_name.upper():
                        no_invoice = n_val

    # Fallback pencarian teks jika belum ditemukan di tabel
    if not no_invoice:
        inv_match = re.search(r'(?:INVOICE|FAKTUR\s*PAJAK)\s+([A-Za-z0-9\/\.-]+)\s+[0-9]{2}[-/\.][0-9]{2}[-/\.][0-9]{4}', full_text, re.IGNORECASE)
        if not inv_match:
            inv_match = re.search(r'(?:13\.\s*Faktur\s*Pajak|11\.\s*Packing\s*List|13\.\s*Invoice)\s*:\s*([A-Za-z0-9\/\.-]+)', full_text, re.IGNORECASE)
        if inv_match:
            no_invoice = inv_match.group(1).strip()
            
    if not no_surat_jalan:
        sj_match = re.search(r'SURAT\s*JALAN\s+([A-Za-z0-9\/\.-]+)\s+[0-9]{2}[-/\.][0-9]{2}[-/\.][0-9]{4}', full_text, re.IGNORECASE)
        if not sj_match:
            sj_match = re.search(r'(?:Surat\s*Jalan|Delivery\s*Order|DO)\s*:\s*([A-Za-z0-9\/\.-]+)', full_text, re.IGNORECASE)
        if sj_match:
            no_surat_jalan = sj_match.group(1).strip()
            
    # 5. Data Pengemas (Jenis Kemasan)
    jenis_kemasan = "UNPACKAGE"
    kemasan_match = re.search(r'22\.\s*Jenis\s*Kemasan\s*:\s*([^\n\r]+)', full_text, re.IGNORECASE)
    if kemasan_match:
        raw_kem = kemasan_match.group(1).strip()
        jenis_kemasan = re.split(r'\d+\.\s*Jumlah\s*Kemasan', raw_kem, flags=re.IGNORECASE)[0].strip()
        
    # 6. Total Bruto & Netto Dokumen
    doc_bruto = 0.0
    doc_netto = 0.0
    bruto_match = re.search(r'25\.\s*Berat\s*Kotor\s*([\d\.,]+)', full_text, re.IGNORECASE)
    if bruto_match:
        doc_bruto = _parse_float(bruto_match.group(1))
        
    netto_match = re.search(r'26\.\s*Berat\s*Bersih\s*(?:\(Kg\))?\s*([\d\.,]+)', full_text, re.IGNORECASE)
    if netto_match:
        doc_netto = _parse_float(netto_match.group(1))

    # 7. Ekstraksi Item Barang dari Tabel PDF
    items = []
    
    for tbl in all_tables:
        for r in tbl:
            if not r:
                continue
            # Deteksi baris barang: Kolom 0 angka seri, Kolom 1 ada 'Pos Tarif'
            if len(r) >= 5 and r[0] and str(r[0]).strip().isdigit() and r[1] and 'Pos Tarif' in str(r[1]):
                seri = str(r[0]).strip()
                uraian_cell = str(r[1])
                js_cell = str(r[3])
                harga_cell = str(r[4])
                
                # Pos Tarif / HS
                m_hs = re.search(r'Pos\s*Tarif/HS\s*:\s*([0-9\.]+)', uraian_cell)
                kode_hs = m_hs.group(1).strip() if m_hs else ""
                
                # Kode Barang
                m_kd = re.search(r'Kode\s*Barang\s*:\s*([A-Za-z0-9\._-]+)', uraian_cell)
                kode_barang = m_kd.group(1).strip() if m_kd else ""
                
                # Nama Barang (bersihkan label Merk, Tipe, Ukuran)
                m_nm = re.search(r'Kode\s*Barang\s*:\s*[^\r\n]+[\r\n]+-\s*([^\r\n]+)', uraian_cell)
                raw_nm = m_nm.group(1).strip() if m_nm else ""
                nama_barang = re.split(r',\s*Merk\s*:', raw_nm, flags=re.IGNORECASE)[0].strip()
                
                # Jumlah & Satuan & Netto dari Kolom 29
                # Format khas: "- 1.0000 \n PCE (PIECE) \n - 1.1770 \n -"
                js_lines = [l.strip() for l in js_cell.split('\n') if l.strip()]
                jumlah = 1.0
                satuan = "PCE (PIECE)"
                netto_item = doc_netto
                
                num_count = 0
                for line in js_lines:
                    m_num = re.match(r'^-\s*([\d\.,]+)$', line)
                    if m_num:
                        val = _parse_float(m_num.group(1))
                        if num_count == 0:
                            jumlah = val
                            num_count += 1
                        elif num_count == 1:
                            netto_item = val
                            num_count += 1
                    elif '(' in line or any(u in line for u in ['PCE', 'KGM', 'SET', 'ROL', 'UNIT', 'BOX', 'CAN', 'BTL', 'BAG', 'PCS']):
                        clean_line = line.replace('- 0', '').replace('-', '').strip()
                        if clean_line:
                            satuan = clean_line
                            
                # Nilai Barang dari Kolom 30
                # Format khas: "- 14,263.0000 \n - 0"
                nilai_item = 0.0
                h_lines = [l.strip() for l in harga_cell.split('\n') if l.strip()]
                for l in h_lines:
                    m_h = re.match(r'^-\s*([\d,]+(?:\.\d+)?|[\d\.]+(?:,\d+)?)$', l)
                    if m_h:
                        val = _parse_float(m_h.group(1))
                        if val > 0:
                            nilai_item = val
                            break
                if nilai_item == 0 and h_lines:
                    m_any = re.search(r'([\d,]+(?:\.\d+)?)', h_lines[0])
                    if m_any:
                        nilai_item = _parse_float(m_any.group(1))
                        
                harga_satuan = round(nilai_item / jumlah, 4) if jumlah > 0 else nilai_item
                
                # Hitung proporsional bruto
                if doc_netto > 0 and doc_bruto > 0:
                    bruto_item = round((netto_item / doc_netto) * doc_bruto, 4)
                else:
                    bruto_item = doc_bruto if len(items) == 0 else netto_item
                    
                items.append({
                    'seri': seri,
                    'kode_barang': kode_barang,
                    'nama_barang': nama_barang,
                    'kode_hs': kode_hs,
                    'kemasan': jenis_kemasan,
                    'jumlah': jumlah,
                    'satuan': satuan,
                    'bruto': bruto_item,
                    'netto': netto_item,
                    'harga_satuan': harga_satuan,
                    'nilai_barang': nilai_item
                })

    # Fallback jika struktur tabel tidak terbaca (misal PDF rasterized)
    if not items:
        hs_match = re.search(r'Pos\s*Tarif/HS\s*:\s*([0-9\.]+)', full_text, re.IGNORECASE)
        kd_match = re.search(r'Kode\s*Barang\s*:\s*([A-Za-z0-9\._-]+)', full_text, re.IGNORECASE)
        nama_match = re.search(r'Kode\s*Barang\s*:\s*[^\r\n]+[\r\n]+-\s*([^\r\n]+)', full_text, re.IGNORECASE)
        harga_penyerahan_match = re.search(r'18\.\s*Harga\s*Penyerahan\s*:\s*([\d\.,]+)', full_text, re.IGNORECASE)
        
        kode_hs_val = hs_match.group(1).strip() if hs_match else ""
        kode_brg_val = kd_match.group(1).strip() if kd_match else ""
        raw_nm = nama_match.group(1).strip() if nama_match else "Barang Pemasukan BC 4.0"
        nama_brg_val = re.split(r',\s*Merk\s*:', raw_nm, flags=re.IGNORECASE)[0].strip()
        nilai_penyerahan = _parse_float(harga_penyerahan_match.group(1)) if harga_penyerahan_match else 0.0
        
        items.append({
            'seri': '1',
            'kode_barang': kode_brg_val,
            'nama_barang': nama_brg_val,
            'kode_hs': kode_hs_val,
            'kemasan': jenis_kemasan,
            'jumlah': 1.0,
            'satuan': 'PCE (PIECE)',
            'bruto': doc_bruto,
            'netto': doc_netto,
            'harga_satuan': nilai_penyerahan,
            'nilai_barang': nilai_penyerahan
        })

    return {
        'dokumen': dokumen,
        'nomor_aju': nomor_aju,
        'nomor_daftar': nomor_daftar,
        'tanggal_daftar': tanggal_daftar,
        'nama_pemasok': nama_pemasok,
        'no_surat_jalan': no_surat_jalan,
        'no_invoice': no_invoice,
        'items': items,
        'item_list': items
    }

def _parse_float(val_str):
    """Helper untuk membersihkan dan mengonversi format angka lokal/internasional ke float"""
    if not val_str:
        return 0.0
    s = str(val_str).replace(' ', '').replace('$', '').replace('Rp', '').replace('IDR', '')
    if '.' in s and ',' in s:
        if s.find('.') < s.find(','):
            # Format Indonesia: 1.234,56
            s = s.replace('.', '').replace(',', '.')
        else:
            # Format Internasional: 1,234.56
            s = s.replace(',', '')
    elif ',' in s:
        # 1234,56 -> 1234.56
        s = s.replace(',', '.')
    try:
        return float(s)
    except:
        return 0.0

def generate_bc40_excel(doc_data_list, output_path: str):
    """
    Menulis data hasil ekstraksi ke dalam template master DATABASE LOKAL BC 4.0.xlsx.
    """
    if os.path.exists(TEMPLATE_PATH):
        wb = openpyxl.load_workbook(TEMPLATE_PATH)
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "DATABASE LOKAL BC 4.0"
        headers = ['NO', 'DOKUMEN', 'NOMOR AJU', 'NOMOR DAFTAR', 'TANGGAL DAFTAR', 
                   'NAMA PEMASOK/PENGIRIM', 'NO SURAT JALAN', 'NO INVOICE', 'SERI', 
                   'KODE BARANG', 'NAMA BARANG', 'KODE HS', 'KEMASAN', 'JUMLAH', 
                   'SATUAN', 'BRUTO', 'NETTO', 'HARGA SATUAN', 'NILAI BARANG']
        ws.append(['KAWASAN BERIKAT PT ZINUS DREAM INDONESIA'])
        ws.append(['LAPORAN PEMASUKAN BARANG PER DOKUMEN PABEAN'])
        ws.append(['PERIODE'])
        ws.append(headers)

    ws = wb.active
    start_row = 5
    
    # Hapus row placeholder lama jika ada
    max_r = ws.max_row
    if max_r >= 12:
        ws.delete_rows(5, max_r - 4)

    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    regular_font = Font(name='Calibri', size=10)
    bold_font = Font(name='Calibri', size=10, bold=True)
    
    current_row = start_row
    no_counter = 1
    
    # Urutkan dokumen secara konsisten berdasarkan Nomor Aju (Ascending)
    sorted_docs = sorted(
        doc_data_list,
        key=lambda d: str(d.get('nomor_aju', '')).strip()
    )
    
    for doc in sorted_docs:
        # Urutkan item barang berdasarkan nomor seri (1, 2, 3...)
        items = sorted(
            doc.get('items', []),
            key=lambda it: int(it.get('seri', 0)) if str(it.get('seri', '')).isdigit() else str(it.get('seri', ''))
        )
        for it in items:
            ws.cell(row=current_row, column=1, value=no_counter).alignment = Alignment(horizontal='center', vertical='center')
            ws.cell(row=current_row, column=2, value=doc.get('dokumen', 'BC 4.0')).alignment = Alignment(horizontal='center', vertical='center')
            ws.cell(row=current_row, column=3, value=doc.get('nomor_aju', '')).alignment = Alignment(horizontal='center', vertical='center')
            ws.cell(row=current_row, column=4, value=doc.get('nomor_daftar', '')).alignment = Alignment(horizontal='center', vertical='center')
            ws.cell(row=current_row, column=5, value=doc.get('tanggal_daftar', '')).alignment = Alignment(horizontal='center', vertical='center')
            ws.cell(row=current_row, column=6, value=doc.get('nama_pemasok', '')).alignment = Alignment(horizontal='left', vertical='center')
            ws.cell(row=current_row, column=7, value=doc.get('no_surat_jalan', '')).alignment = Alignment(horizontal='center', vertical='center')
            ws.cell(row=current_row, column=8, value=doc.get('no_invoice', '')).alignment = Alignment(horizontal='center', vertical='center')
            ws.cell(row=current_row, column=9, value=it.get('seri', '1')).alignment = Alignment(horizontal='center', vertical='center')
            ws.cell(row=current_row, column=10, value=it.get('kode_barang', '')).alignment = Alignment(horizontal='center', vertical='center')
            ws.cell(row=current_row, column=11, value=it.get('nama_barang', '')).alignment = Alignment(horizontal='left', vertical='center')
            ws.cell(row=current_row, column=12, value=it.get('kode_hs', '')).alignment = Alignment(horizontal='center', vertical='center')
            ws.cell(row=current_row, column=13, value=it.get('kemasan', '')).alignment = Alignment(horizontal='center', vertical='center')
            
            # Jumlah
            c_jml = ws.cell(row=current_row, column=14, value=it.get('jumlah', 0))
            c_jml.number_format = '#,##0.0000' if isinstance(it.get('jumlah'), float) and not it.get('jumlah').is_integer() else '#,##0'
            c_jml.alignment = Alignment(horizontal='right', vertical='center')
            
            # Satuan (Contoh: PCE (PIECE))
            ws.cell(row=current_row, column=15, value=it.get('satuan', 'PCE (PIECE)')).alignment = Alignment(horizontal='center', vertical='center')
            
            # Bruto (4 desimal)
            c_bruto = ws.cell(row=current_row, column=16, value=it.get('bruto', 0))
            c_bruto.number_format = '#,##0.0000'
            c_bruto.alignment = Alignment(horizontal='right', vertical='center')
            
            # Netto (4 desimal)
            c_netto = ws.cell(row=current_row, column=17, value=it.get('netto', 0))
            c_netto.number_format = '#,##0.0000'
            c_netto.alignment = Alignment(horizontal='right', vertical='center')
            
            # Harga Satuan (4 desimal)
            c_hrg = ws.cell(row=current_row, column=18, value=it.get('harga_satuan', 0))
            c_hrg.number_format = '#,##0.0000'
            c_hrg.alignment = Alignment(horizontal='right', vertical='center')
            
            # Nilai Barang (4 desimal)
            c_nilai = ws.cell(row=current_row, column=19, value=it.get('nilai_barang', 0))
            c_nilai.number_format = '#,##0.0000'
            c_nilai.alignment = Alignment(horizontal='right', vertical='center')
            
            for col in range(1, 20):
                cell = ws.cell(row=current_row, column=col)
                cell.font = regular_font
                cell.border = thin_border
                
            current_row += 1
            no_counter += 1

    # Total Summary Row
    end_data_row = current_row - 1
    if end_data_row >= start_row:
        total_row = current_row
        c_tot_no = ws.cell(row=total_row, column=1, value=f"=COUNTA(A{start_row}:A{end_data_row})")
        c_tot_no.alignment = Alignment(horizontal='center', vertical='center')
        
        c_tot_bruto = ws.cell(row=total_row, column=16, value=f"=SUM(P{start_row}:P{end_data_row})")
        c_tot_bruto.number_format = '#,##0.0000'
        c_tot_bruto.alignment = Alignment(horizontal='right', vertical='center')
        
        c_tot_netto = ws.cell(row=total_row, column=17, value=f"=SUM(Q{start_row}:Q{end_data_row})")
        c_tot_netto.number_format = '#,##0.0000'
        c_tot_netto.alignment = Alignment(horizontal='right', vertical='center')
        
        c_tot_nilai = ws.cell(row=total_row, column=19, value=f"=SUM(S{start_row}:S{end_data_row})")
        c_tot_nilai.number_format = '#,##0.0000'
        c_tot_nilai.alignment = Alignment(horizontal='right', vertical='center')
        
        for col in range(1, 20):
            cell = ws.cell(row=total_row, column=col)
            cell.font = bold_font
            cell.border = thin_border
            cell.fill = PatternFill(start_color='F2F4F7', end_color='F2F4F7', fill_type='solid')

    wb.save(output_path)
    return output_path
