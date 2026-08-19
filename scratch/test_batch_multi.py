import os, sys
sys.path.insert(0, r'd:/new project/export_tools_app')
sys.path.insert(0, r'd:/new project')
import openpyxl

from komparasi import process_mass_reconciliation

peb_files = [
    r'd:/new project/export_tools_app/uploads/609_ID2608-8023.pdf',
    r'd:/new project/export_tools_app/uploads/616_ID2608-8113.pdf'
]

cipl_files = [
    r'd:/new project/export_tools_app/uploads/ID2608-8023_ZN_IDN_MBLMEDUKU519287_ETD_08.21.2026_BKAMZ992N18413JK0_AMZ_ML.pdf',
    r'd:/new project/export_tools_app/uploads/ID2608-8113_SUB_INVOICE_0993631007_F3.pdf'
]

out_xlsx = r'd:/new project/scratch/test_batch_multi.xlsx'

res = process_mass_reconciliation(peb_files, cipl_files, out_xlsx)
print("Reconciliation Summary:", res.get("reconciliation_summary"))

wb = openpyxl.load_workbook(out_xlsx, data_only=True)
for sheetname in wb.sheetnames:
    if sheetname == "Dashboard Summary": continue
    ws = wb[sheetname]
    print("\n" + "="*80)
    print("SHEET:", sheetname)
    print("="*80)
    for r in range(3, 21):
        no = ws.cell(r, 1).value
        elem = ws.cell(r, 2).value
        peb = ws.cell(r, 3).value
        cipl = ws.cell(r, 4).value
        st = ws.cell(r, 5).value
        if elem in ["Nama Penerima (Consignee)", "Alamat penerima", "Nama Pembeli (Buyer)", "Alamat Pembeli"]:
            print(f"[{no}] {elem:30s} | PEB : {repr(peb)}")
            print(f"     {'':30s} | CIPL: {repr(cipl)} | Status: {st}")
