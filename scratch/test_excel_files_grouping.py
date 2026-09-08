import sys
import os
from collections import defaultdict
sys.path.insert(0, r'd:/new project')
sys.path.insert(0, r'd:/new project/export_tools_app')

from app.services.draf_peb_service import (
    generate_draf_peb_excel,
    create_zip_bundle
)
import openpyxl

doc1 = {
    'filename': 'CIPL_1.pdf',
    'invoice_no': 'ID2608-8141',
    'invoice_date': '2026-08-14',
    'packing_list_date': '2026-08-14',
    'date_8digit': '20260830',
    'seq_6digit': '000635',
    'nomor_aju': '000030ZIK48020260830000635',
    'excel_group': '1',
    'country': 'US',
    'pelabuhan_tujuan': 'USCHS',
    'etd': '2026-08-21',
    'bl_no': 'ONEYJKTG65320401',
    'bl_date': '2026-08-21',
    'vessel': 'SINAR CARITA',
    'voy_no': '026N',
    'flag_code': 'SG',
    'buyer_name': 'ZINUS INC.',
    'consignee_name': 'WALMART, INC.',
    'consignee_address': 'BENTONVILLE AR',
    'total_qty': 100,
    'total_fob': 5000.0,
    'total_nw': 1000.0,
    'total_gw': 1133.0,
    'gross_weight': 1133.0,
    'net_weight': 1000.0,
    'bruto': 1133.0,
    'netto': 1000.0,
    'total_volume': 10.5,
    'items': [
        {
            'code': '94042120',
            'des': 'MATTRESS ITEM 1',
            'po': '1004113825',
            'sku': 'ZU-MFMA1OZI-10Q',
            'f_code': 'F.MFM.08Q.000.W',
            'qt': 100,
            'price': 50.0,
            'fob': 5000.0,
            'nw': 1000.0,
            'gw': 1133.0
        }
    ]
}

doc2 = {
    'filename': 'CIPL_2.pdf',
    'invoice_no': 'ID2608-8142',
    'invoice_date': '2026-08-15',
    'packing_list_date': '2026-08-15',
    'date_8digit': '20260830',
    'seq_6digit': '000636',
    'nomor_aju': '000030ZIK48020260830000636',
    'excel_group': '1',
    'country': 'US',
    'pelabuhan_tujuan': 'USCHS',
    'etd': '2026-08-21',
    'bl_no': 'ONEYJKTG65320401',
    'bl_date': '2026-08-21',
    'vessel': 'SINAR CARITA',
    'voy_no': '026N',
    'flag_code': 'SG',
    'buyer_name': 'ZINUS INC.',
    'consignee_name': 'WALMART, INC.',
    'consignee_address': 'BENTONVILLE AR',
    'total_qty': 50,
    'total_fob': 2500.0,
    'total_nw': 500.0,
    'total_gw': 566.5,
    'gross_weight': 566.5,
    'net_weight': 500.0,
    'bruto': 566.5,
    'netto': 500.0,
    'total_volume': 5.2,
    'items': [
        {
            'code': '94042120',
            'des': 'MATTRESS ITEM 2',
            'po': '1004113826',
            'sku': 'ZU-MFMA1OZI-08F',
            'f_code': 'F.MFM.08F.000.W',
            'qt': 50,
            'price': 50.0,
            'fob': 2500.0,
            'nw': 500.0,
            'gw': 566.5
        }
    ]
}

doc3 = {
    'filename': 'CIPL_3.pdf',
    'invoice_no': 'ID2608-8143',
    'invoice_date': '2026-08-16',
    'packing_list_date': '2026-08-16',
    'date_8digit': '20260830',
    'seq_6digit': '000637',
    'nomor_aju': '000030ZIK48020260830000637',
    'excel_group': '2',
    'country': 'US',
    'pelabuhan_tujuan': 'USCHS',
    'etd': '2026-08-21',
    'bl_no': 'ONEYJKTG65320402',
    'bl_date': '2026-08-21',
    'vessel': 'SINAR CARITA',
    'voy_no': '026N',
    'flag_code': 'SG',
    'buyer_name': 'ZINUS INC.',
    'consignee_name': 'WALMART, INC.',
    'consignee_address': 'BENTONVILLE AR',
    'total_qty': 200,
    'total_fob': 10000.0,
    'total_nw': 2000.0,
    'total_gw': 2266.0,
    'gross_weight': 2266.0,
    'net_weight': 2000.0,
    'bruto': 2266.0,
    'netto': 2000.0,
    'total_volume': 21.0,
    'items': [
        {
            'code': '94042120',
            'des': 'MATTRESS ITEM 3',
            'po': '1004113827',
            'sku': 'ZU-MFMA1OZI-12K',
            'f_code': 'F.MFM.12K.000.W',
            'qt': 200,
            'price': 50.0,
            'fob': 10000.0,
            'nw': 2000.0,
            'gw': 2266.0
        }
    ]
}

documents = [doc1, doc2, doc3]
print(f"Total documents: {len(documents)}")

global_data = {
    'kode_kantor': '040300',
    'kode_kantor_periksa': '150300',
    'kode_kantor_ekspor': '040300',
    'pelabuhan_muat': 'IDTPP',
    'pelabuhan_ekspor': 'IDTPP',
    'tanggal_periksa': '2026-08-18',
    'kota_pernyataan': 'TANGERANG',
    'tanggal_pernyataan': '2026-08-30',
    'nama_pernyataan': 'EUN SUN KANG',
    'jabatan_pernyataan': 'MANAGER',
    'ndpbm_kurs': 17960,
    'kode_daerah_asal': '3603',
    'kode_negara_asal': 'ID',
    'kode_jenis_ekspor': '1',
    'statement_perbedaan_harga': 'T'
}

excel_groups = defaultdict(list)
for doc in documents:
    egrp = str(doc.get('excel_group', '1')).strip() or '1'
    excel_groups[egrp].append(doc)

generated_paths = []
generated_files = []
template_path = r'D:\DOC\contoh draf peb.xlsx'

for g_idx, (egrp_id, g_docs) in enumerate(excel_groups.items()):
    out_excel_path = generate_draf_peb_excel(g_docs, global_data, template_path=template_path)
    fname = os.path.basename(out_excel_path)
    generated_paths.append(out_excel_path)
    generated_files.append({
        'filename': fname,
        'doc_count': len(g_docs),
        'documents': g_docs
    })
    print(f"Generated File #{g_idx+1}: {fname} ({len(g_docs)} No Aju)")

assert len(generated_files) == 2, "Expected 2 Excel files"

# Verify File 1 (contains doc1 & doc2)
wb1 = openpyxl.load_workbook(generated_paths[0], data_only=True)
ws1_h = wb1['HEADER']
print("\n--- File 1 HEADER rows ---")
aju1 = ws1_h.cell(row=2, column=1).value
aju2 = ws1_h.cell(row=3, column=1).value
print(f"File 1 Row 2 Aju: {aju1}")
print(f"File 1 Row 3 Aju: {aju2}")
assert aju1 == '000030ZIK48020260830000635'
assert aju2 == '000030ZIK48020260830000636'

# Verify File 2 (contains doc3)
wb2 = openpyxl.load_workbook(generated_paths[1], data_only=True)
ws2_h = wb2['HEADER']
print("\n--- File 2 HEADER rows ---")
aju3 = ws2_h.cell(row=2, column=1).value
print(f"File 2 Row 2 Aju: {aju3}")
assert aju3 == '000030ZIK48020260830000637'

# Verify ZIP bundle
zip_path = create_zip_bundle(generated_paths)
print(f"\nCreated ZIP Bundle: {zip_path}")
assert os.path.exists(zip_path) and os.path.getsize(zip_path) > 0

print("\nALL EXCEL GROUPING VERIFICATIONS PASSED 100% PERFECTLY!")
