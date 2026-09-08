811 EXCELLENCE DRIVE BENTONVILLE AR 72716 UNITED STATESimport sys
import os
sys.path.insert(0, r'd:/new project')
sys.path.insert(0, r'd:/new project/export_tools_app')

from app.services.draf_peb_service import (
    consolidate_cipl_documents,
    generate_draf_peb_excel
)
import openpyxl

doc1 = {
    'filename': 'CIPL_A.pdf',
    'invoice_no': 'ID2608-8141',
    'invoice_date': '2026-08-14',
    'packing_list_date': '2026-08-14',
    'date_8digit': '20260830',
    'seq_6digit': '000635',
    'nomor_aju': '000030ZIK48020260830000635',
    'group_id': '1',
    'country': 'US',
    'pelabuhan_tujuan': 'USCHS',
    'etd': '2026-08-21',
    'bl_no': 'ONEYJKTG65320401',
    'bl_date': '2026-08-21',
    'vessel': 'SINAR CARITA',
    'voy_no': '026N',
    'flag_code': 'SG',
    'buyer_name': 'ZINUS INC.',
    'consignee_name': 'WALMART, INC.',
    'consignee_address': 'BENTONVILLE AR',
    'total_qty': 100,
    'total_fob': 5000.0,
    'total_nw': 1000.0,
    'total_gw': 1133.0,
    'gross_weight': 1133.0,
    'net_weight': 1000.0,
    'bruto': 1133.0,
    'netto': 1000.0,
    'total_volume': 10.5,
    'items': [
        {
            'code': '94042120',
            'des': 'MATTRESS ITEM 1',
            'po': '1004113825',
            'sku': 'ZU-MFMA1OZI-10Q',
            'f_code': 'F.MFM.08Q.000.W',
            'qt': 100,
            'price': 50.0,
            'fob': 5000.0,
            'nw': 1000.0,
            'gw': 1133.0
        }
    ]
}

doc2 = {
    'filename': 'CIPL_B.pdf',
    'invoice_no': 'ID2608-8142',
    'invoice_date': '2026-08-15',
    'packing_list_date': '2026-08-15',
    'date_8digit': '20260830',
    'seq_6digit': '000636',
    'nomor_aju': '000030ZIK48020260830000636',
    'group_id': '1',
    'country': 'US',
    'pelabuhan_tujuan': 'USCHS',
    'etd': '2026-08-21',
    'bl_no': 'ONEYJKTG65320401',
    'bl_date': '2026-08-21',
    'vessel': 'SINAR CARITA',
    'voy_no': '026N',
    'flag_code': 'SG',
    'buyer_name': 'ZINUS INC.',
    'consignee_name': 'WALMART, INC.',
    'consignee_address': 'BENTONVILLE AR',
    'total_qty': 50,
    'total_fob': 2500.0,
    'total_nw': 500.0,
    'total_gw': 566.5,
    'gross_weight': 566.5,
    'net_weight': 500.0,
    'bruto': 566.5,
    'netto': 500.0,
    'total_volume': 5.2,
    'items': [
        {
            'code': '94042120',
            'des': 'MATTRESS ITEM 2',
            'po': '1004113826',
            'sku': 'ZU-MFMA1OZI-08F',
            'f_code': 'F.MFM.08F.000.W',
            'qt': 50,
            'price': 50.0,
            'fob': 2500.0,
            'nw': 500.0,
            'gw': 566.5
        }
    ]
}

doc3 = {
    'filename': 'CIPL_C.pdf',
    'invoice_no': 'ID2608-8143',
    'invoice_date': '2026-08-16',
    'packing_list_date': '2026-08-16',
    'date_8digit': '20260830',
    'seq_6digit': '000637',
    'nomor_aju': '000030ZIK48020260830000637',
    'group_id': '2',
    'country': 'US',
    'pelabuhan_tujuan': 'USCHS',
    'etd': '2026-08-21',
    'bl_no': 'ONEYJKTG65320402',
    'bl_date': '2026-08-21',
    'vessel': 'SINAR CARITA',
    'voy_no': '026N',
    'flag_code': 'SG',
    'buyer_name': 'ZINUS INC.',
    'consignee_name': 'WALMART, INC.',
    'consignee_address': 'BENTONVILLE AR',
    'total_qty': 200,
    'total_fob': 10000.0,
    'total_nw': 2000.0,
    'total_gw': 2266.0,
    'gross_weight': 2266.0,
    'net_weight': 2000.0,
    'bruto': 2266.0,
    'netto': 2000.0,
    'total_volume': 21.0,
    'items': [
        {
            'code': '94042120',
            'des': 'MATTRESS ITEM 3',
            'po': '1004113827',
            'sku': 'ZU-MFMA1OZI-12K',
            'f_code': 'F.MFM.12K.000.W',
            'qt': 200,
            'price': 50.0,
            'fob': 10000.0,
            'nw': 2000.0,
            'gw': 2266.0
        }
    ]
}

raw_docs = [doc1, doc2, doc3]
print(f"Original uploaded docs: {len(raw_docs)}")

consolidated = consolidate_cipl_documents(raw_docs, start_seq=635)
print(f"Consolidated Aju count: {len(consolidated)}")

assert len(consolidated) == 2, f"Expected 2 Aju, got {len(consolidated)}"
assert consolidated[0]['total_qty'] == 150, f"Expected 150, got {consolidated[0]['total_qty']}"
assert consolidated[0]['total_fob'] == 7500.0, f"Expected 7500.0, got {consolidated[0]['total_fob']}"
assert len(consolidated[0]['items']) == 2, f"Expected 2 items, got {len(consolidated[0]['items'])}"
assert len(consolidated[0]['sub_docs']) == 2, f"Expected 2 sub_docs, got {len(consolidated[0]['sub_docs'])}"

assert consolidated[1]['total_qty'] == 200, f"Expected 200, got {consolidated[1]['total_qty']}"
assert consolidated[1]['total_fob'] == 10000.0, f"Expected 10000.0, got {consolidated[1]['total_fob']}"
assert len(consolidated[1]['items']) == 1, f"Expected 1 item, got {len(consolidated[1]['items'])}"

global_data = {
    'kode_kantor': '040300',
    'kode_kantor_periksa': '150300',
    'kode_kantor_ekspor': '040300',
    'pelabuhan_muat': 'IDTPP',
    'pelabuhan_ekspor': 'IDTPP',
    'tanggal_periksa': '2026-08-18',
    'kota_pernyataan': 'TANGERANG',
    'tanggal_pernyataan': '2026-08-30',
    'nama_pernyataan': 'EUN SUN KANG',
    'jabatan_pernyataan': 'MANAGER',
    'ndpbm_kurs': 17960,
    'kode_daerah_asal': '3603',
    'kode_negara_asal': 'ID',
    'kode_jenis_ekspor': '1',
    'statement_perbedaan_harga': 'T'
}

out_excel = generate_draf_peb_excel(consolidated, global_data, template_path=r'D:\DOC\contoh draf peb.xlsx')
print(f"Generated Excel: {out_excel}")

wb = openpyxl.load_workbook(out_excel, data_only=True)
ws_h = wb['HEADER']
print("\n--- SHEET HEADER ---")
for r in range(2, 4):
    aju = ws_h.cell(row=r, column=1).value
    fob = ws_h.cell(row=r, column=69).value
    ndpbm = ws_h.cell(row=r, column=75).value
    print(f"Row {r}: No Aju = {aju} | FOB = {fob} | NDPBM = {ndpbm}")

assert ws_h.cell(row=2, column=69).value == 7500.0, "HEADER Row 2 FOB should be 7500.0"
assert ws_h.cell(row=3, column=69).value == 10000.0, "HEADER Row 3 FOB should be 10000.0"

ws_d = wb['DOKUMEN']
print("\n--- SHEET DOKUMEN ---")
for r in range(1, ws_d.max_row + 1):
    row_vals = [ws_d.cell(row=r, column=c).value for c in range(1, 6)]
    if any(row_vals):
        print(f"Row {r}: {row_vals}")

ws_b = wb['BARANG']
print("\n--- SHEET BARANG ---")
for r in range(2, ws_b.max_row + 1):
    row_vals = [ws_b.cell(row=r, column=1).value, ws_b.cell(row=r, column=2).value, ws_b.cell(row=r, column=5).value, ws_b.cell(row=r, column=11).value, ws_b.cell(row=r, column=29).value]
    print(f"Barang Row {r}: {row_vals}")

print("\nALL VERIFICATIONS PASSED PERFECTLY!")
