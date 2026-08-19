import os, sys
sys.path.insert(0, r'd:/new project/export_tools_app')
sys.path.insert(0, r'd:/new project')
import openpyxl

from komparasi import process_mass_reconciliation

peb_files = [
    r'd:/new project/export_tools_app/uploads/636_ID2608-8142.pdf'
]

cipl_files = [
    r'd:/new project/export_tools_app/uploads/ID2608-8142_SUB_INVOICE_0993631159_F3.pdf'
]

out_xlsx = r'd:/new project/scratch/test_8142.xlsx'

res = process_mass_reconciliation(peb_files, cipl_files, out_xlsx)
print("Reconciliation Summary:", res.get("reconciliation_summary"))

wb = openpyxl.load_workbook(out_xlsx, data_only=True)
for sheetname in wb.sheetnames:
    if sheetname == "Dashboard Summary": continue
    ws = wb[sheetname]
    print("\n" + "="*100)
    print("SHEET:", sheetname)
    print("="*100)
    print("DETAIL COMPARISON TABLE:")
    for r in range(21, 35):
        row_vals = [ws.cell(r, c).value for c in range(1, 15)]
        if any(row_vals):
            print(f"Row {r:2d}: {row_vals}")
