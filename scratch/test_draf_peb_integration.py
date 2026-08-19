import os
import sys
import openpyxl

sys.path.insert(0, r"d:/new project/export_tools_app")
from app.services.draf_peb_service import parse_cipl_file_for_draf_peb, generate_draf_peb_excel

def test_full_draf_peb():
    pdf_path = r"d:/new project/export_tools_app/uploads/ID2607-7849_SUB_INVOICE_0993631151_F3.pdf"
    template_path = r"D:\DOC\contoh draf peb.xlsx"

    print("1. Testing CIPL PDF parsing...")
    cipl_data = parse_cipl_file_for_draf_peb(pdf_path)
    assert cipl_data["invoice_no"] == "ID2607-7849", f"Expected ID2607-7849, got {cipl_data['invoice_no']}"
    assert cipl_data["total_qty"] == 693, f"Expected 693, got {cipl_data['total_qty']}"
    print(f"   [OK] CIPL parsed: Inv={cipl_data['invoice_no']}, Qty={cipl_data['total_qty']}, FOB=${cipl_data['total_fob']}")

    print("2. Testing Excel generation...")
    form_data = {
        "nomor_aju": "000030ZIK48020260814000777",
        "kode_kantor": "040300",
        "kode_kantor_periksa": "150300",
        "kode_kantor_ekspor": "040300",
        "kode_negara_tujuan": "US",
        "pelabuhan_muat": "IDTPP",
        "pelabuhan_tujuan": "USCHS",
        "pelabuhan_ekspor": "IDTPP",
        "tanggal_ekspor": "2026-08-28",
        "tanggal_periksa": "2026-08-25",
        "kota_pernyataan": "TANGERANG",
        "tanggal_pernyataan": "2026-08-16",
        "nama_pernyataan": "EUN SUN KANG",
        "jabatan_pernyataan": "MANAGER",
        "ndpbm_kurs": 17960,
        "fob": cipl_data["total_fob"],
        "bruto": cipl_data["total_gw"],
        "netto": cipl_data["total_nw"],
        "total_qty": cipl_data["total_qty"],
        "nomor_pengangkut": "026N",
        "kode_bendera": "SG"
    }

    out_file = generate_draf_peb_excel(cipl_data, form_data, template_path=template_path)
    assert os.path.exists(out_file), f"Output file does not exist: {out_file}"
    print(f"   [OK] Excel generated at: {out_file}")

    print("3. Validating sheets and cell contents...")
    wb = openpyxl.load_workbook(out_file)
    assert len(wb.sheetnames) == 21, f"Expected 21 sheets, got {len(wb.sheetnames)}"
    
    # Check HEADER
    ws_h = wb["HEADER"]
    assert ws_h["A2"].value == "000030ZIK48020260814000777"
    assert ws_h["C2"].value == "040300"
    assert ws_h["BQ2"].value == cipl_data["total_fob"]
    assert ws_h["CC2"].value == cipl_data["total_nw"]
    print("   [OK] HEADER values verified")

    # Check ENTITAS
    ws_e = wb["ENTITAS"]
    assert ws_e["A2"].value == "000030ZIK48020260814000777"
    print("   [OK] ENTITAS values verified")

    # Check DOKUMEN
    ws_d = wb["DOKUMEN"]
    assert ws_d["A2"].value == "000030ZIK48020260814000777"
    print("   [OK] DOKUMEN values verified")

    # Check BARANG
    ws_b = wb["BARANG"]
    assert ws_b["A2"].value == "000030ZIK48020260814000777"
    assert ws_b["C2"].value == "94042120"
    assert ws_b["K2"].value == 693
    print("   [OK] BARANG values verified")

    print("\nALL INTEGRATION TESTS PASSED 100%!")

if __name__ == "__main__":
    test_full_draf_peb()
