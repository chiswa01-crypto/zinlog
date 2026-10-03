import sys
import os

root_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
if root_dir not in sys.path:
    sys.path.insert(0, root_dir)

app_dir = os.path.join(root_dir, 'export_tools_app')
if app_dir not in sys.path:
    sys.path.insert(0, app_dir)

from export_tools_app.app.services.draf_peb_service import generate_draf_peb_excel, create_zip_bundle
import openpyxl

doc_a = {
    "nomor_aju": "000030ZIK48020261003000635",
    "invoice_no": "INV-A-101",
    "invoice_date": "2026-10-01",
    "country": "US",
    "consignee_name": "BUYER A CORP",
    "consignee_address": "100 A STREET, NY",
    "total_qty": 100,
    "total_fob": 5000.0,
    "total_gw": 1200.0,
    "total_nw": 1000.0,
    "total_volume": 10.0,
    "items": [{
        "code": "94042120",
        "f_code": "ITEM-A-01",
        "des": "MATTRESS A",
        "po": "PO-A",
        "sku": "SKU-A",
        "qt": 100,
        "nw": 1000.0,
        "fob": 5000.0,
        "price": 50.0
    }]
}

doc_b = {
    "nomor_aju": "000030ZIK48020261003000636",
    "invoice_no": "INV-B-202",
    "invoice_date": "2026-10-02",
    "country": "US",
    "consignee_name": "BUYER B CORP",
    "consignee_address": "200 B STREET, LA",
    "total_qty": 200,
    "total_fob": 10000.0,
    "total_gw": 2400.0,
    "total_nw": 2000.0,
    "total_volume": 20.0,
    "items": [{
        "code": "94042120",
        "f_code": "ITEM-B-02",
        "des": "MATTRESS B",
        "po": "PO-B",
        "sku": "SKU-B",
        "qt": 200,
        "nw": 2000.0,
        "fob": 10000.0,
        "price": 50.0
    }]
}

doc_c = {
    "nomor_aju": "000030ZIK48020261003000637",
    "invoice_no": "INV-C-303",
    "invoice_date": "2026-10-03",
    "country": "US",
    "consignee_name": "BUYER C CORP",
    "consignee_address": "300 C STREET, CHI",
    "total_qty": 300,
    "total_fob": 15000.0,
    "total_gw": 3600.0,
    "total_nw": 3000.0,
    "total_volume": 30.0,
    "items": [{
        "code": "94042120",
        "f_code": "ITEM-C-03",
        "des": "MATTRESS C",
        "po": "PO-C",
        "sku": "SKU-C",
        "qt": 300,
        "nw": 3000.0,
        "fob": 15000.0,
        "price": 50.0
    }]
}

global_data = {
    "kode_kantor": "040300",
    "kode_kantor_periksa": "150300",
    "kode_kantor_ekspor": "040300",
    "tanggal_periksa": "2026-10-03",
    "tanggal_pernyataan": "2026-10-03",
    "nama_pernyataan": "TEST USER",
    "jabatan_pernyataan": "MANAGER"
}

out_dir = os.path.join(root_dir, 'scratch', 'test_output')
os.makedirs(out_dir, exist_ok=True)

file1 = generate_draf_peb_excel([doc_a], global_data, output_path=os.path.join(out_dir, 'Draf_File1_INV-A.xlsx'))
file2 = generate_draf_peb_excel([doc_b], global_data, output_path=os.path.join(out_dir, 'Draf_File2_INV-B.xlsx'))
file3 = generate_draf_peb_excel([doc_c], global_data, output_path=os.path.join(out_dir, 'Draf_File3_INV-C.xlsx'))

wb1 = openpyxl.load_workbook(file1)
wb2 = openpyxl.load_workbook(file2)
wb3 = openpyxl.load_workbook(file3)

inv1 = wb1["DOKUMEN"].cell(row=4, column=4).value # 380 invoice
inv2 = wb2["DOKUMEN"].cell(row=4, column=4).value
inv3 = wb3["DOKUMEN"].cell(row=4, column=4).value

print("File 1 Invoice:", inv1)
print("File 2 Invoice:", inv2)
print("File 3 Invoice:", inv3)

des1 = wb1["BARANG"].cell(row=2, column=5).value
des2 = wb2["BARANG"].cell(row=2, column=5).value
des3 = wb3["BARANG"].cell(row=2, column=5).value

print("File 1 Item Description:", des1)
print("File 2 Item Description:", des2)
print("File 3 Item Description:", des3)

assert inv1 == "INV-A-101", f"Expected INV-A-101, got {inv1}"
assert inv2 == "INV-B-202", f"Expected INV-B-202, got {inv2}"
assert inv3 == "INV-C-303", f"Expected INV-C-303, got {inv3}"

assert des1 == "MATTRESS A", f"Expected MATTRESS A, got {des1}"
assert des2 == "MATTRESS B", f"Expected MATTRESS B, got {des2}"
assert des3 == "MATTRESS C", f"Expected MATTRESS C, got {des3}"

print("SUCCESS! All 3 files are 100% separate and distinct!")
