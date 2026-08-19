import os, sys
sys.path.insert(0, r'd:/new project/export_tools_app')
sys.path.insert(0, r'd:/new project')
import openpyxl

from komparasi import process_mass_reconciliation

pdf_peb = r'd:/new project/export_tools_app/uploads/616_ID2608-8113.pdf'
pdf_cipl = r'd:/new project/export_tools_app/uploads/ID2608-8113_SUB_INVOICE_0993631007_F3.pdf'
out_xlsx = r'd:/new project/scratch/test_out_8113.xlsx'

result = process_mass_reconciliation([pdf_peb], [pdf_cipl], out_xlsx)
print("Result:", result)

wb = openpyxl.load_workbook(out_xlsx, data_only=True)
ws = wb.active
print("Sheet title:", ws.title)
for r in range(22, 27):
    row_vals = [ws.cell(r, c).value for c in range(1, 23)]
    print(f"Row {r}: {row_vals}")

print("\nCIPL Cell Values & Formats:")
for c in range(13, 23):
    cell = ws.cell(24, c)
    print(f"  Col {c} ({ws.cell(23, c).value}): value={repr(cell.value)} (type={type(cell.value).__name__}), format={cell.number_format}")
