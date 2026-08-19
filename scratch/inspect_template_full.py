import openpyxl
from openpyxl.utils import get_column_letter

wb = openpyxl.load_workbook(r'D:\DOC\data real.xlsx', data_only=False)
ws = wb['bulan']

print("=== MERGED RANGES ===")
for m in ws.merged_cells.ranges:
    print(m)

print("\n=== ALL NON-EMPTY CELLS ===")
for r in range(1, ws.max_row + 1):
    for c in range(1, min(ws.max_column + 1, 30)):
        cell = ws.cell(row=r, column=c)
        if cell.value is not None:
            col_letter = get_column_letter(c)
            fill_hex = cell.fill.start_color.rgb if cell.fill and cell.fill.start_color else None
            font_desc = f"{cell.font.name} {cell.font.size}pt bold={cell.font.bold}" if cell.font else "NoFont"
            border_info = f"top={cell.border.top.style if cell.border and cell.border.top else None}, bot={cell.border.bottom.style if cell.border and cell.border.bottom else None}"
            align_desc = f"h={cell.alignment.horizontal}, v={cell.alignment.vertical}" if cell.alignment else "NoAlign"
            print(f"{col_letter}{r}: val={cell.value!r} | fill={fill_hex} | font={font_desc} | num_fmt={cell.number_format!r} | {border_info} | {align_desc}")

print("\n=== ROW DIMENSIONS ===")
for r in range(1, ws.max_row + 1):
    rd = ws.row_dimensions[r]
    if rd.height is not None or rd.hidden:
        print(f"Row {r}: height={rd.height}, hidden={rd.hidden}")

print("\n=== COLUMN WIDTHS ===")
for col_name, cd in ws.column_dimensions.items():
    if cd.width is not None:
        print(f"Col {col_name}: width={cd.width}")
