import os, sys
sys.path.insert(0, r'd:/new project/export_tools_app')
sys.path.insert(0, r'd:/new project')
import openpyxl

from komparasi import process_mass_reconciliation

pdf_peb = r'd:/new project/export_tools_app/uploads/609_ID2608-8023.pdf'
pdf_cipl = r'd:/new project/export_tools_app/uploads/ID2608-8023_ZN_IDN_MBLMEDUKU519287_ETD_08.21.2026_BKAMZ992N18413JK0_AMZ_ML.pdf'
out_xlsx = r'd:/new project/scratch/test_out_8023.xlsx'

result = process_mass_reconciliation([pdf_peb], [pdf_cipl], out_xlsx)
print("Result summary:", result.get('reconciliation_summary'))

wb = openpyxl.load_workbook(out_xlsx, data_only=True)
ws = wb['Master_ID2608-8023']
print("Using sheet:", ws.title)
for r in range(24, 29):
    p_row = [ws.cell(r, c).value for c in range(1, 12)]
    c_row = [ws.cell(r, c).value for c in range(13, 23)]
    status = ws.cell(r, 2).value
    print(f"Row {r-23}: Status={status}")
    print(f"  PEB : {p_row}")
    print(f"  CIPL: {c_row}")
